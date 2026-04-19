import random
import pygame
import sys
from Constants import *
from Field import Field
from Character import Character
from Cursor import *
import numpy as np
import AI
from GameMaster import GameMaster
import GameMaster as GM
from MapData import MapData
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
from pathlib import Path

AI_SELECTION_LABELS = (
    'Player Input',
    'Baseline AI',
    'Random AI',
    'Personality Cores AI',
    'Aggressive Personality Cores AI',
    'Strategic Personality Cores AI',
    'Survival Personality Cores AI',
    'Kill One By One AI',
    'Disable AI'
)



# === Board placement (top-left of the 640x640 grid) ===
# Lower this to move the whole board (and its A-H / 1-8 labels) higher on screen.
BOARD_POS_X = WIDTH // 2 - 320
BOARD_POS_Y = 40

# === Game Log placement (anchored under the board) ===
LOG_W = WIDTH - 40              # full-ish width (tweak as you like)
LOG_H = 270                     # panel height (was 280)
LOG_X = 20                      # left margin
LOG_Y = BOARD_POS_Y + 700

# --- Cream UI colors for the bottom log panel (like your screenshot) ---
UI_PANEL  = (251, 247, 242)   # panel fill
UI_HEADER = (233, 226, 214)   # header strip
UI_BORDER = (30, 30, 30)      # dark border
UI_TEXT   = (20, 20, 20)      # text

LOG_COLOR_GAME    = (20, 20, 20)      # black
LOG_COLOR_ROUND   = (255, 165, 0)     # orange
LOG_COLOR_SUMMARY = (22, 138, 36)     # green
LOG_COLOR_P1      = (54, 92, 168)     # blue
LOG_COLOR_P2      = (178, 64, 64)     # red

# --- Scrollbar colors ---
SB_TRACK       = (220, 213, 200)
SB_THUMB       = (160, 150, 135)
SB_THUMB_HOVER = (145, 135, 120)
SB_THUMB_DRAG  = (130, 120, 105)


class GameMainExportMixin:
    def export_game_log(self) -> None:
        """Export all played games into one workbook with per-game log + summary sheets."""

        def convert_to_board_position(row, col) -> str:
            try:
                row_int = int(row)
                col_int = int(col)
            except (TypeError, ValueError):
                return ""
            if col_int < 0:
                return ""
            column_letter = chr(ord('A') + col_int)
            return f"{column_letter}{row_int}"

        def normalize_position(value) -> str:
            if value is None:
                return ""
            value_str = str(value).strip()
            if not value_str:
                return ""

            upper_value = value_str.upper()
            if len(upper_value) >= 2 and upper_value[0].isalpha() and upper_value[1:].isdigit():
                return upper_value

            if "," in value_str:
                parts = [part.strip() for part in value_str.split(",", 1)]
                if len(parts) == 2:
                    return convert_to_board_position(parts[0], parts[1])

            return value_str

        def make_sheet_title(raw_title: str, used_titles: set[str]) -> str:
            invalid_chars = set('[]:*?/\\')
            cleaned = ''.join('_' if ch in invalid_chars else ch for ch in raw_title).strip() or 'Sheet'
            cleaned = cleaned[:31]
            if cleaned not in used_titles:
                used_titles.add(cleaned)
                return cleaned

            suffix = 2
            while True:
                suffix_text = f"_{suffix}"
                base = cleaned[:31 - len(suffix_text)]
                candidate = f"{base}{suffix_text}"
                if candidate not in used_titles:
                    used_titles.add(candidate)
                    return candidate
                suffix += 1

        def as_int(value, default=0) -> int:
            try:
                return int(value)
            except (TypeError, ValueError):
                return default

        def as_float(value, default=0.0) -> float:
            try:
                return float(value)
            except (TypeError, ValueError):
                return default

        def normalize_map_value(value) -> str:
            if value is None:
                return ""
            map_value = str(value).strip()
            if not map_value:
                return ""
            if map_value.lower().startswith("map "):
                return map_value[4:].strip()
            return map_value

        def sanitize_for_excel(value):
            if value is None:
                return ""
            s = str(value)
            if not s:
                return ""
            # Prevent Excel from interpreting leading characters as formulas or causing removals
            if s[0] in ('=', '+', '-', '@'):
                return "'" + s
            return s

        try:
            if not self.game_log:
                print("No game log data to export.")
                return

            wb = openpyxl.Workbook()
            wb.remove(wb.active)

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')
            )
            summary_title_font = Font(bold=True, color="FFFFFF", size=13)
            summary_title_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            summary_metric_font = Font(bold=True)
            summary_metric_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
            summary_value_font = Font(bold=False, color="000000")
            summary_value_fill = PatternFill(fill_type=None)

            # Ensure logs are exported oldest -> newest (top -> down in sheets).
            # `_game_log_archive` stores every entry in chronological order (no cap).
            archive = getattr(self, '_game_log_archive', None)
            if archive:
                ordered_logs = [entry for entry in archive if isinstance(entry, dict)]
            else:
                # Fallback: use the capped UI log (reversed to chronological)
                ordered_logs = [entry for entry in reversed(self.game_log) if isinstance(entry, dict)]
            if not ordered_logs:
                print("No structured game log data to export.")
                return

            games_data: dict[int, dict] = {}

            def get_or_create_game(game_no: int) -> dict:
                if game_no not in games_data:
                    games_data[game_no] = {
                        "rows": [],
                        "match": 0,
                        "map": "",
                        "winner": "",
                        "duration": 0.0,
                        "damage_p1": 0,
                        "damage_p2": 0,
                        "kills_p1": 0,
                        "kills_p2": 0,
                        "obj_p1": 0,
                        "obj_p2": 0,
                        "p1_model": "",
                        "p2_model": "",
                        "duration_sum": 0.0,
                    }
                return games_data[game_no]

            for entry in ordered_logs:
                game_no = as_int(entry.get("game", 1), 1)
                game_data = get_or_create_game(game_no)
                game_data["rows"].append(entry)

                game_data["match"] = max(game_data["match"], as_int(entry.get("match", 0), 0))
                game_data["duration"] = max(game_data["duration"], as_float(entry.get("duration", entry.get("time", 0.0)), 0.0))
                game_data["duration_sum"] += as_float(entry.get("time", 0.0), 0.0)

                if not game_data["map"] and entry.get("map"):
                    game_data["map"] = str(entry.get("map", ""))
                if entry.get("winner"):
                    game_data["winner"] = str(entry.get("winner", ""))
                if not game_data["p1_model"] and entry.get("p1_model"):
                    game_data["p1_model"] = str(entry.get("p1_model", ""))
                if not game_data["p2_model"] and entry.get("p2_model"):
                    game_data["p2_model"] = str(entry.get("p2_model", ""))

                team = as_int(entry.get("team", 0), 0)
                damage = as_int(entry.get("damage", 0), 0)
                if team == 1:
                    game_data["damage_p1"] += damage
                elif team == 2:
                    game_data["damage_p2"] += damage

                if str(entry.get("action_type", "")).lower() == "ko":
                    if team == 1:
                        game_data["kills_p2"] += 1
                    elif team == 2:
                        game_data["kills_p1"] += 1

                game_data["obj_p1"] = max(game_data["obj_p1"], as_int(entry.get("objective_control_p1", 0), 0))
                game_data["obj_p2"] = max(game_data["obj_p2"], as_int(entry.get("objective_control_p2", 0), 0))

            all_game_keys = sorted(games_data.keys())
            # Export all logged games. This keeps export resilient when UI limits are
            # edited with free-form numeric input (including values > 10).
            export_game_keys = all_game_keys

            # Filter out games that have no player/action rows (these are often empty
            # placeholder games created by session logic like a logged game_start)
            def game_has_actions(gdata: dict) -> bool:
                for e in gdata.get('rows', []):
                    team = int(e.get('team', 0) or 0)
                    action_type = str(e.get('action_type', '') or '').lower()
                    damage = as_int(e.get('damage', 0), 0)
                    heal = as_int(e.get('heal', 0), 0)
                    cls = str(e.get('class', '') or '').strip()
                    if team in (1, 2):
                        return True
                    if action_type in ('attack', 'move', 'heal', 'pass', 'ko'):
                        return True
                    if damage > 0 or heal > 0:
                        return True
                    if cls:
                        return True
                return False

            export_game_keys = [k for k in export_game_keys if game_has_actions(games_data.get(k, {}))]

            total_games = len(export_game_keys)
            if total_games == 0:
                print("No game log data to export (no games with action rows).")
                return

            # Derive overall game wins from per-game logged match results to be robust
            p1_wins = 0
            p2_wins = 0
            for k in export_game_keys:
                gd = games_data.get(k, {})
                # prefer explicit per-game winner when available
                game_winner = str(gd.get('winner', '') or '').strip().upper()
                if game_winner:
                    if 'P1' in game_winner or game_winner.endswith('1') or 'PLAYER 1' in game_winner:
                        p1_wins += 1
                        continue
                    if 'P2' in game_winner or game_winner.endswith('2') or 'PLAYER 2' in game_winner:
                        p2_wins += 1
                        continue

                # fallback: count any logged winner fields in this game's rows
                p1_matches = 0
                p2_matches = 0
                for e in gd.get('rows', []):
                    w_raw = str(e.get('winner', '') or '').strip().upper()
                    if not w_raw:
                        continue
                    if 'P1' in w_raw or w_raw.endswith('1') or 'PLAYER 1' in w_raw:
                        p1_matches += 1
                    elif 'P2' in w_raw or w_raw.endswith('2') or 'PLAYER 2' in w_raw:
                        p2_matches += 1
                if p1_matches > p2_matches:
                    p1_wins += 1
                elif p2_matches > p1_matches:
                    p2_wins += 1
            # winrate as fraction (0.66 = 66%) for proper Excel formatting
            winrate_p1 = (p1_wins / total_games) if total_games else 0.0
            winrate_p2 = (p2_wins / total_games) if total_games else 0.0

            headers = [
                "No.", "Log Type", "Log", "Match", "Round", "Team", "Time (seconds)", "Class", "Health",
                "Position", "Action type", "Action name", "Target Class", "Target Position",
                "Damage", "Heal", "Target Health before", "Target Health after"
            ]
            # column widths must match headers count
            column_widths = [6, 12, 36, 8, 8, 6, 14, 18, 10, 18, 14, 18, 18, 18, 10, 10, 20, 20]
            used_titles: set[str] = set()

            # helper to split prefixed log tags (e.g. "MATCH : ...") for CSV clarity
            def split_log_tag(raw_log: str) -> tuple[str, str]:
                if raw_log is None:
                    return "", ""
                log_str = str(raw_log).strip()
                if not log_str:
                    return "", ""
                # try splitting at first ':' to extract tag
                parts = log_str.split(':', 1)
                if len(parts) == 2:
                    left = parts[0].strip()
                    right = parts[1].strip()
                    # treat left as tag when it's short and alphabetic-ish
                    compact = left.replace(' ', '')
                    if 0 < len(compact) <= 12 and compact.isalpha():
                        return compact.upper(), right
                return "", log_str

            # csv_rows removed; we write cleaned log directly into XLSX

            # Use the filtered export_game_keys for summary metrics
            first_game_key = min(export_game_keys)
            first_game_data = games_data[first_game_key]
            overall_duration = round(sum(games_data[k]["duration_sum"] for k in export_game_keys), 2)

            overall_summary_ws = wb.create_sheet(make_sheet_title("Overall_summary", used_titles), index=0)
            overall_summary_ws.column_dimensions['A'].width = 28
            overall_summary_ws.column_dimensions['B'].width = 32

            overall_title_cell = overall_summary_ws.cell(row=1, column=1, value="Overall Summary")
            overall_title_cell.font = summary_title_font
            overall_title_cell.fill = PatternFill(start_color="4B0082", end_color="4B0082", fill_type="solid")
            overall_title_cell.alignment = header_alignment
            overall_title_cell.border = thin_border
            overall_summary_ws.merge_cells('A1:B1')
            overall_summary_ws.cell(row=1, column=2).border = thin_border

            overall_summary_data = [
                ("P1 Model", first_game_data["p1_model"] or "Unknown"),
                ("P2 Model", first_game_data["p2_model"] or "Unknown"),
                ("Configured Game Limit", max(1, as_int(getattr(self, "game_limit", 1), 1))),
                ("Configured Match Limit", max(1, as_int(getattr(self, "match_limit", 1), 1))),
                ("Game", total_games),
                ("P1 Wins", p1_wins),
                ("P2 Wins", p2_wins),
                ("Duration (seconds)", overall_duration),
                ("Win rate P1", winrate_p1),
                ("Win rate P2", winrate_p2),
            ]

            for row_idx, (metric, value) in enumerate(overall_summary_data, start=2):
                metric_cell = overall_summary_ws.cell(row=row_idx, column=1, value=metric)
                metric_cell.font = summary_metric_font
                metric_cell.fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
                metric_cell.alignment = Alignment(horizontal="left", vertical="center")
                metric_cell.border = thin_border

                value_cell = overall_summary_ws.cell(row=row_idx, column=2, value=value)
                value_cell.font = summary_value_font
                value_cell.fill = summary_value_fill
                value_cell.alignment = Alignment(horizontal="left", vertical="center")
                value_cell.border = thin_border
                # format winrate cells as percentage
                if isinstance(value, float) and 'Win rate' in metric:
                    value_cell.number_format = '0.00%'

            for game_no in sorted(export_game_keys):
                game_data = games_data[game_no]

                log_ws = wb.create_sheet(make_sheet_title(f"G{game_no}_game_log", used_titles))
                for i, width in enumerate(column_widths, 1):
                    log_ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

                for col, header in enumerate(headers, 1):
                    cell = log_ws.cell(row=1, column=col, value=header)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                    cell.border = thin_border

                for index, entry in enumerate(game_data["rows"], start=1):
                    position = normalize_position(entry.get("position", ""))
                    target_position = normalize_position(entry.get("target_position", ""))
                    # split log tag into type + cleaned message, then write both into XLSX
                    raw_log = entry.get("log", "")
                    log_type, clean_log = split_log_tag(raw_log)
                    row_values = [
                        index,
                        log_type,
                        clean_log,
                        as_int(entry.get("match", 0), 0),
                        as_int(entry.get("round", 0), 0),
                        as_int(entry.get("team", 0), 0),
                        as_float(entry.get("time", 0.0), 0.0),
                        str(entry.get("class", "")),
                        as_int(entry.get("health", 0), 0),
                        position,
                        str(entry.get("action_type", "")),
                        str(entry.get("action_name", "")),
                        str(entry.get("target_class", "")),
                        target_position,
                        as_int(entry.get("damage", 0), 0),
                        as_int(entry.get("heal", 0), 0),
                        as_int(entry.get("target_hp_before", 0), 0),
                        as_int(entry.get("target_hp_after", 0), 0),
                    ]
                    # sanitize string cells to avoid Excel treating them as formulas
                    sanitized = [sanitize_for_excel(v) if isinstance(v, str) else v for v in row_values]
                    log_ws.append(sanitized)
                    for col in range(1, len(row_values) + 1):
                        log_ws.cell(row=index + 1, column=col).border = thin_border

                summary_ws = wb.create_sheet(make_sheet_title(f"G{game_no}_summary", used_titles))
                summary_ws.column_dimensions['A'].width = 28
                summary_ws.column_dimensions['B'].width = 32

                title_cell = summary_ws.cell(row=1, column=1, value=f"Game {game_no} Summary")
                title_cell.font = summary_title_font
                title_cell.fill = summary_title_fill
                title_cell.alignment = header_alignment
                title_cell.border = thin_border
                summary_ws.merge_cells('A1:B1')
                summary_ws.cell(row=1, column=2).border = thin_border

                total_obj = game_data["obj_p1"] + game_data["obj_p2"]
                obj_p1_pct = (game_data["obj_p1"] / total_obj * 100.0) if total_obj else 0.0
                obj_p2_pct = (game_data["obj_p2"] / total_obj * 100.0) if total_obj else 0.0

                # compute per-game match wins by inspecting any logged winner fields
                p1_matches_in_game = 0
                p2_matches_in_game = 0
                for e in game_data.get('rows', []):
                    winner = str(e.get('winner', '') or '').strip().upper()
                    if not winner:
                        continue
                    if 'P1' in winner or winner.endswith('1') or 'PLAYER 1' in winner:
                        p1_matches_in_game += 1
                    elif 'P2' in winner or winner.endswith('2') or 'PLAYER 2' in winner:
                        p2_matches_in_game += 1

                total_matches_in_game = p1_matches_in_game + p2_matches_in_game
                winrate_p1_game = (p1_matches_in_game / total_matches_in_game) if total_matches_in_game else 0.0
                winrate_p2_game = (p2_matches_in_game / total_matches_in_game) if total_matches_in_game else 0.0

                summary_data = [
                    ("Map", normalize_map_value(game_data["map"]) or "Unknown"),
                    ("Winner", game_data["winner"] or "Unknown"),
                    ("Matches P1", p1_matches_in_game),
                    ("Matches P2", p2_matches_in_game),
                    ("Win rate P1", winrate_p1_game),
                    ("Win rate P2", winrate_p2_game),
                    ("Duration (seconds)", round(game_data["duration_sum"], 2)),
                    ("Damage Dealt P1", game_data["damage_p1"]),
                    ("Damage Dealt P2", game_data["damage_p2"]),
                    ("Kills by P1", game_data["kills_p1"]),
                    ("Kills by P2", game_data["kills_p2"]),
                    ("Objective Control P1", f"{game_data['obj_p1']} ({obj_p1_pct:.2f}%)"),
                    ("Objective Control P2", f"{game_data['obj_p2']} ({obj_p2_pct:.2f}%)"),
                ]

                for row_idx, (metric, value) in enumerate(summary_data, start=2):
                    metric_cell = summary_ws.cell(row=row_idx, column=1, value=metric)
                    metric_cell.font = summary_metric_font
                    metric_cell.fill = summary_metric_fill
                    metric_cell.alignment = Alignment(horizontal="left", vertical="center")
                    metric_cell.border = thin_border

                    value_cell = summary_ws.cell(row=row_idx, column=2, value=value)
                    value_cell.font = summary_value_font
                    value_cell.fill = summary_value_fill
                    value_cell.alignment = Alignment(horizontal="left", vertical="center")
                    value_cell.border = thin_border

            desktop = Path.home() / "Desktop"
            export_dir = desktop / "RPG Simulation Export"
            export_dir.mkdir(parents=True, exist_ok=True)

            filename = f"game_log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            filepath = export_dir / filename
            wb.save(filepath)
            print(f"Game log exported to: {filepath}")
            # No extra CSV file is created; XLSX now contains `Log Type` and cleaned `Log` columns

        except Exception as e:
            print(f"Error exporting game log: {e}")
            import traceback
            traceback.print_exc()
